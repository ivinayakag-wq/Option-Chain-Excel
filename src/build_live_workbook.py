import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .build_workbook import (
    AMBER,
    CHAIN_COLUMNS,
    CHAIN_NUMBER_FORMATS,
    FONT_NAME,
    GREEN,
    LIGHT_GREEN,
    LIGHT_RED,
    LIGHT_SLATE,
    RED,
    SLATE,
    SUMMARY_COLUMNS,
    SUMMARY_NUMBER_FORMATS,
    _format_header,
    _rgb,
)
from .config import ROOT
from .upstox import INDICES

LIVE_PATH = ROOT / "OptionChain Live.xlsx"
SNAPSHOT_PATH = ROOT / "OptionChain.xlsx"
LAST_ROW = 25
SUMMARY_LAST_ROW = 9

READ_ME = [
    ("title", "Option Chain Live — Read Me"),
    ("sub", "NIFTY · BANKNIFTY · FINNIFTY · MIDCPNIFTY · SENSEX"),
    ("gap", ""),
    ("head", "What this is"),
    ("body", "A pushed live view. A Python service streams Upstox market data over a WebSocket and writes into this workbook every half second while the market is open. You do not press Refresh All anywhere."),
    ("body", "Cells change by themselves. The line in A2 on every sheet is the heartbeat: LIVE, the local time, and how old the newest tick is."),
    ("gap", ""),
    ("head", "Before the market"),
    ("body", "1. Log in once per trading day: open http://127.0.0.1:8765/login and approve. The token dies at 03:30 IST every morning."),
    ("body", "2. Double-click 'Start Live Option Chain.bat'. It opens this workbook if it is not already open and starts writing."),
    ("body", "Leave the black window open. Closing it stops the feed."),
    ("gap", ""),
    ("head", "What updates, and how fast"),
    ("body", "The 20 strikes around at-the-money on each sheet, plus the strike itself, refresh every 0.5 seconds: LTP, change, OI, DOI, volume, IV, bid and ask. Anything outside that window is not shown."),
    ("body", "The window follows the market. If the index drifts far enough that ATM moves out of the middle, the rows shift by themselves; the new edge strikes fill in within a second or two."),
    ("body", "Summary refreshes every 60 seconds from the full chain, because PCR and max pain are only meaningful across every strike. Its Age column tells you how old those numbers are."),
    ("gap", ""),
    ("head", "Reading a chain sheet"),
    ("body", "Calls on the left of the Strike column, puts on the right. OI is open interest, DOI is the change in OI versus the previous session, IV is implied volatility in percent, LTP is last traded price, Chg% is versus the previous close. Shaded cells are in-the-money. Amber rows are the max-pain strike."),
    ("gap", ""),
    ("head", "If something looks wrong"),
    ("body", "A2 turns amber — the feed has gone quiet for more than five seconds. It usually recovers on its own within a few seconds."),
    ("body", "A2 turns red — the token expired or the broker refused the connection. Open the login URL and approve, then restart the launcher."),
    ("body", "Nothing updates and Excel feels stuck — close this workbook and run the launcher again."),
    ("gap", ""),
    ("head", "Notes"),
    ("body", "This file is written by a program. Typing into the chain sheets will be overwritten within half a second, and Excel's undo history is cleared on every write. Edit in the snapshot workbook instead."),
    ("body", "The snapshot workbook, OptionChain.xlsx, still works the old way: press Data ▸ Refresh All. Use it if the live service is down."),
    ("body", "Data is licensed to you personally through your Upstox account. Do not redistribute the feed."),
]


def build_read_me(wb: Workbook) -> None:
    ws = wb.create_sheet("Read Me")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 118
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 2
    row = 1
    for kind, text in READ_ME:
        cell = ws.cell(row=row, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            cell.font = Font(name=FONT_NAME, size=20, bold=True, color="0F172A")
            ws.row_dimensions[row].height = 30
        elif kind == "sub":
            cell.font = Font(name=FONT_NAME, size=11, color="64748B")
        elif kind == "head":
            cell.font = Font(name=FONT_NAME, size=12, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E293B")
            ws.row_dimensions[row].height = 22
            row += 1
            continue
        elif kind == "body":
            cell.font = Font(name=FONT_NAME, size=10.5, color="1F2937")
            ws.row_dimensions[row].height = 16 * max(1, (len(text) // 112) + 1)
        else:
            ws.row_dimensions[row].height = 8
        row += 1


def build_shell(path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    build_read_me(wb)

    summary = wb.create_sheet("Summary")
    summary.sheet_properties.tabColor = "334155"
    summary["A1"] = "Summary — nearest expiry, full chain"
    summary["A1"].font = Font(name=FONT_NAME, size=16, bold=True, color="0F172A")
    summary.row_dimensions[1].height = 24
    summary.row_dimensions[3].height = 8

    for symbol in INDICES:
        ws = wb.create_sheet(symbol)
        ws.sheet_properties.tabColor = "1E6B34" if symbol != "SENSEX" else "8B1E1E"
        ws["A1"] = f"{symbol} — live, ATM ± 10 strikes"
        ws["A1"].font = Font(name=FONT_NAME, size=16, bold=True, color="0F172A")
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[3].height = 8

    wb.save(path)


def _write_headers(ws, columns, header_row: int = 4) -> None:
    for index, (_, friendly, _) in enumerate(columns, start=1):
        ws.Cells(header_row, index).Value = friendly


def _apply_chain_formats(ws, app, summary_row: int) -> None:
    _write_headers(ws, CHAIN_COLUMNS)
    _format_header(
        ws,
        21,
        4,
        {
            **{c: SLATE for c in (1, 2, 3)},
            **{c: GREEN for c in range(4, 12)},
            **{12: (15, 23, 42)},
            **{c: RED for c in range(13, 21)},
            **{21: SLATE},
        },
    )
    for col in range(1, 22):
        rng = ws.Range(f"{ws.Cells(5, col).Address}:{ws.Cells(LAST_ROW, col).Address}")
        rng.Font.Name = FONT_NAME
        rng.Font.Size = 10
        if col in CHAIN_NUMBER_FORMATS:
            rng.NumberFormat = CHAIN_NUMBER_FORMATS[col]

    widths = {
        "A": 11, "B": 11, "C": 10, "D": 12, "E": 11, "F": 11, "G": 8, "H": 10,
        "I": 9, "J": 10, "K": 10, "L": 10, "M": 10, "N": 10, "O": 9, "P": 10,
        "Q": 8, "R": 11, "S": 11, "T": 12, "U": 6,
    }
    for letter, width in widths.items():
        ws.Columns(letter).ColumnWidth = width

    ws.Columns("B").NumberFormat = "@"
    ws.Range(f"L5:L{LAST_ROW}").Font.Bold = True
    ws.Range(f"L5:L{LAST_ROW}").HorizontalAlignment = -4108
    ws.Range(f"L5:L{LAST_ROW}").Interior.Color = _rgb(*LIGHT_SLATE)

    ce_block = ws.Range(f"D5:K{LAST_ROW}")
    rule = ce_block.FormatConditions.Add(2, 1, "=$C5>$L5")
    rule.Interior.Color = _rgb(*LIGHT_GREEN)

    pe_block = ws.Range(f"M5:T{LAST_ROW}")
    rule = pe_block.FormatConditions.Add(2, 1, "=$C5<$L5")
    rule.Interior.Color = _rgb(*LIGHT_RED)

    whole = ws.Range(f"A5:U{LAST_ROW}")
    rule = whole.FormatConditions.Add(2, 1, f"=$L5=Summary!$G${summary_row}")
    rule.Interior.Color = _rgb(*AMBER)

    for col, tint in (("D", (46, 125, 50)), ("T", (139, 30, 30))):
        try:
            bar = ws.Range(f"{col}5:{col}{LAST_ROW}").FormatConditions.AddDatabar()
            bar.BarColor.Color = _rgb(*tint)
            bar.ShowValue = True
        except Exception:
            pass

    app.ActiveWindow.FreezePanes = False
    ws.Activate()
    ws.Range("D5").Select()
    app.ActiveWindow.FreezePanes = True


def _apply_summary_formats(ws) -> None:
    _write_headers(ws, SUMMARY_COLUMNS)
    _format_header(ws, 18, 4, {c: SLATE for c in range(1, 19)})
    ws.Columns("B").NumberFormat = "@"
    for col in range(1, 19):
        rng = ws.Range(
            f"{ws.Cells(5, col).Address}:{ws.Cells(SUMMARY_LAST_ROW, col).Address}"
        )
        rng.Font.Name = FONT_NAME
        rng.Font.Size = 10
        if col in SUMMARY_NUMBER_FORMATS:
            rng.NumberFormat = SUMMARY_NUMBER_FORMATS[col]
    ws.Range(f"E5:F{SUMMARY_LAST_ROW}").Font.Bold = True
    ws.Range(f"G5:I{SUMMARY_LAST_ROW}").Font.Bold = True
    ws.Range(f"G5:G{SUMMARY_LAST_ROW}").Font.Color = _rgb(*RED)

    widths = {
        "A": 12, "B": 11, "C": 11, "D": 10, "E": 10, "F": 10, "G": 10, "H": 11,
        "I": 13, "J": 11, "K": 11, "L": 13, "M": 13, "N": 11, "O": 11, "P": 8,
        "Q": 7, "R": 8,
    }
    for letter, width in widths.items():
        ws.Columns(letter).ColumnWidth = width

    rule = ws.Range(f"A5:R{SUMMARY_LAST_ROW}").FormatConditions.Add(2, 1, "=$Q5=TRUE")
    rule.Interior.Color = _rgb(*AMBER)

    for col, tint in (("L", (46, 125, 50)), ("M", (139, 30, 30))):
        try:
            bar = ws.Range(f"{col}5:{col}{SUMMARY_LAST_ROW}").FormatConditions.AddDatabar()
            bar.BarColor.Color = _rgb(*tint)
            bar.ShowValue = True
        except Exception:
            pass


def build_via_excel(path: Path) -> None:
    import win32com.client as win32

    app = win32.DispatchEx("Excel.Application")
    app.Visible = False
    app.DisplayAlerts = False
    wb = None
    try:
        wb = app.Workbooks.Open(str(path))
        for index, symbol in enumerate(INDICES, start=1):
            _apply_chain_formats(wb.Worksheets(symbol), app, summary_row=4 + index)
            wb.Worksheets(symbol).Range("A2").Value = "WAITING — start the live service"
        _apply_summary_formats(wb.Worksheets("Summary"))
        wb.Worksheets("Summary").Range("A2").Value = "WAITING — start the live service"
        wb.Worksheets("Read Me").Activate()
        wb.Save()
    finally:
        if wb is not None:
            wb.Close(SaveChanges=False)
        app.Quit()


def verify(path: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    return [f"{ws.title:<14} tables: {', '.join(ws.tables.keys()) or '-'}" for ws in wb.worksheets]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the live OptionChain workbook")
    parser.add_argument("--path", default=str(LIVE_PATH))
    args = parser.parse_args(argv)
    path = Path(args.path)
    build_shell(path)
    print(f"Shell written: {path}")
    build_via_excel(path)
    print(f"Formatted: {path}")
    for line in verify(path):
        print("  " + line)
    return 0


if __name__ == "__main__":
    sys.exit(main())
